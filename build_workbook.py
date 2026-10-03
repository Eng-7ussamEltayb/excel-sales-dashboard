#!/usr/bin/env python3
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

df = pd.read_csv("data/sales.csv")
df["revenue"] = df["qty"] * df["unit_price"]
df["date"] = pd.to_datetime(df["date"])
df["month"] = df["date"].dt.to_period("M").astype(str)
by_region = df.groupby("region", as_index=False)["revenue"].sum()
by_month = df.groupby("month", as_index=False)["revenue"].sum()

wb = Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(list(df.columns))
for row in df.itertuples(index=False):
    ws.append(list(row))

ws2 = wb.create_sheet("ByRegion")
ws2.append(["region", "revenue"])
for r in by_region.itertuples(index=False):
    ws2.append(list(r))

chart = BarChart()
chart.title = "Revenue by Region"
data = Reference(ws2, min_col=2, min_row=1, max_row=1 + len(by_region))
cats = Reference(ws2, min_col=1, min_row=2, max_row=1 + len(by_region))
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws2.add_chart(chart, "E2")

ws3 = wb.create_sheet("ByMonth")
ws3.append(["month", "revenue"])
for r in by_month.itertuples(index=False):
    ws3.append(list(r))

out = Path("output/sales_dashboard.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Created {out}")
